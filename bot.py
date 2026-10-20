import os
import io
import time
import logging
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application, CommandHandler, MessageHandler,
    CallbackQueryHandler, ContextTypes, filters
)
from PIL import Image, ImageOps
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
import markdown
import csv
import html

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO
)
logger = logging.getLogger(__name__)

TOKEN = os.environ.get("BOT_TOKEN", "").strip()

# ---------- User state (what they want to convert to) ----------
# user_state[uid] = {"mode": "image"|"doc", "target": "png"|"pdf"|...}
USER_STATE = {}

IMAGE_FORMATS = ["PNG", "JPG", "WEBP", "GIF", "BMP", "PDF"]
DOC_FORMATS = ["TXT", "PDF", "DOCX", "MD", "HTML", "CSV"]

# ---------- Menus ----------
def main_menu():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🖼 Image Converter", callback_data="img_menu")],
        [InlineKeyboardButton("📄 Document Converter", callback_data="doc_menu")],
        [InlineKeyboardButton("🔤 Text Case Converter", callback_data="case_menu")],
        [InlineKeyboardButton("ℹ️ About", callback_data="about")],
    ])

def image_menu():
    rows = []
    for fmt in IMAGE_FORMATS:
        rows.append([InlineKeyboardButton(f"➡️ Convert to {fmt}", callback_data=f"img:{fmt}")])
    rows.append([InlineKeyboardButton("⬅️ Back", callback_data="back")])
    return InlineKeyboardMarkup(rows)

def doc_menu():
    rows = []
    for fmt in DOC_FORMATS:
        rows.append([InlineKeyboardButton(f"➡️ Convert to {fmt}", callback_data=f"doc:{fmt}")])
    rows.append([InlineKeyboardButton("⬅️ Back", callback_data="back")])
    return InlineKeyboardMarkup(rows)

def case_menu():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("UPPERCASE", callback_data="case:upper")],
        [InlineKeyboardButton("lowercase", callback_data="case:lower")],
        [InlineKeyboardButton("Title Case", callback_data="case:title")],
        [InlineKeyboardButton("Sentence case", callback_data="case:sentence")],
        [InlineKeyboardButton("Reverse text", callback_data="case:reverse")],
        [InlineKeyboardButton("⬅️ Back", callback_data="back")],
    ])

# ---------- Image conversion ----------
def convert_image(data: bytes, target: str) -> io.BytesIO:
    """Convert image bytes to target format. Returns BytesIO with proper name."""
    img = Image.open(io.BytesIO(data))
    img = ImageOps.exif_transpose(img)  # fix phone rotation

    buf = io.BytesIO()
    target = target.upper()

    if target == "JPG":
        # JPEG doesn't support alpha — flatten to white
        if img.mode in ("RGBA", "LA", "P"):
            bg = Image.new("RGB", img.size, (255, 255, 255))
            if img.mode == "P":
                img = img.convert("RGBA")
            bg.paste(img, mask=img.split()[-1] if img.mode in ("RGBA", "LA") else None)
            img = bg
        else:
            img = img.convert("RGB")
        img.save(buf, "JPEG", quality=92)
        buf.name = "converted.jpg"

    elif target == "PNG":
        if img.mode not in ("RGB", "RGBA", "L", "LA"):
            img = img.convert("RGBA")
        img.save(buf, "PNG", optimize=True)
        buf.name = "converted.png"

    elif target == "WEBP":
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGB")
        img.save(buf, "WEBP", quality=90)
        buf.name = "converted.webp"

    elif target == "GIF":
        img = img.convert("P", palette=Image.ADAPTIVE, colors=256)
        img.save(buf, "GIF")
        buf.name = "converted.gif"

    elif target == "BMP":
        img = img.convert("RGB")
        img.save(buf, "BMP")
        buf.name = "converted.bmp"

    elif target == "PDF":
        img = img.convert("RGB")
        img.save(buf, "PDF", resolution=100.0)
        buf.name = "converted.pdf"

    else:
        raise ValueError(f"Unsupported image format: {target}")

    buf.seek(0)
    return buf

# ---------- Document conversion ----------
def extract_text_from_upload(file_bytes: bytes, filename: str) -> str:
    """Extract text from .txt, .md, .csv, .html, .docx (best-effort)."""
    name = (filename or "").lower()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    # Everything else: try UTF-8, then fallback to latin-1
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")

    if name.endswith(".html") or name.endswith(".htm"):
        # Strip tags roughly
        import re
        text = re.sub(r"<[^>]+>", "", text)
        text = html.unescape(text)

    return text

def text_to_pdf(text: str) -> io.BytesIO:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    margin = 40
    y = height - margin
    line_height = 14
    max_chars = 95

    c.setFont("Helvetica", 10)

    for raw_line in text.split("\n"):
        # Wrap long lines
        line = raw_line
        while len(line) > max_chars:
            c.drawString(margin, y, line[:max_chars])
            line = line[max_chars:]
            y -= line_height
            if y < margin:
                c.showPage()
                c.setFont("Helvetica", 10)
                y = height - margin
        c.drawString(margin, y, line)
        y -= line_height
        if y < margin:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = height - margin

    c.save()
    buf.seek(0)
    buf.name = "converted.pdf"
    return buf

def text_to_docx(text: str) -> io.BytesIO:
    doc = Document()
    for para in text.split("\n"):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    buf.name = "converted.docx"
    return buf

def text_to_markdown(text: str) -> io.BytesIO:
    # Simple: keep text as-is, wrap in md paragraph breaks
    md = "\n\n".join(line for line in text.split("\n"))
    buf = io.BytesIO(md.encode("utf-8"))
    buf.name = "converted.md"
    return buf

def text_to_html(text: str) -> io.BytesIO:
    body = "<br>".join(html.escape(line) for line in text.split("\n"))
    page = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Converted</title></head>
<body><pre>{body}</pre></body></html>"""
    buf = io.BytesIO(page.encode("utf-8"))
    buf.name = "converted.html"
    return buf

def text_to_csv(text: str) -> io.BytesIO:
    buf = io.StringIO()
    writer = csv.writer(buf)
    for line in text.split("\n"):
        writer.writerow([line])
    data = buf.getvalue().encode("utf-8")
    b = io.BytesIO(data)
    b.name = "converted.csv"
    return b

def text_to_txt(text: str) -> io.BytesIO:
    buf = io.BytesIO(text.encode("utf-8"))
    buf.name = "converted.txt"
    return buf

def convert_document(text: str, target: str) -> io.BytesIO:
    target = target.upper()
    if target == "TXT":
        return text_to_txt(text)
    if target == "PDF":
        return text_to_pdf(text)
    if target == "DOCX":
        return text_to_docx(text)
    if target == "MD":
        return text_to_markdown(text)
    if target == "HTML":
        return text_to_html(text)
    if target == "CSV":
        return text_to_csv(text)
    raise ValueError(f"Unsupported document format: {target}")

# ---------- Text case conversion ----------
def apply_case(text: str, mode: str) -> str:
    if mode == "upper":
        return text.upper()
    if mode == "lower":
        return text.lower()
    if mode == "title":
        return text.title()
    if mode == "sentence":
        return ". ".join(s.strip().capitalize() for s in text.split("."))
    if mode == "reverse":
        return text[::-1]
    return text

# ---------- Handlers ----------
async def start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "👋 Welcome to *SB24GZ Official Bot*!\n\n"
        "I convert *images* and *documents* between formats.\n\n"
        "• 🖼 Image → PNG / JPG / WEBP / GIF / BMP / PDF\n"
        "• 📄 Text file → TXT / PDF / DOCX / MD / HTML / CSV\n"
        "• 🔤 Text case → UPPER / lower / Title / reverse\n\n"
        "Pick a tool below 👇",
        parse_mode="Markdown",
        reply_markup=main_menu(),
    )

async def help_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "📖 *How to use:*\n\n"
        "1. Tap a converter button\n"
        "2. Choose the target format\n"
        "3. Send an image, file, or text\n"
        "4. Get the converted file back\n\n"
        "Commands:\n"
        "/start — Main menu\n"
        "/menu — Show menu\n"
        "/cancel — Reset state",
        parse_mode="Markdown",
    )

async def menu_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("Choose a tool:", reply_markup=main_menu())

async def cancel_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    USER_STATE.pop(uid, None)
    await update.message.reply_text("✅ Cancelled. Nothing selected.", reply_markup=main_menu())

async def about(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    text = (
        "🤖 *SB24GZ Official Bot*\n\n"
        "Image & document format converter.\n"
        "Runs entirely on the server — no external APIs.\n"
        "Deployed on Railway 🚂"
    )
    if q:
        await q.answer()
        await q.edit_message_text(text, parse_mode="Markdown", reply_markup=main_menu())
    else:
        await update.message.reply_text(text, parse_mode="Markdown")

async def button_handler(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    uid = update.effective_user.id
    await q.answer()

    if q.data == "back":
        return await q.edit_message_text("Choose a tool:", reply_markup=main_menu())

    if q.data == "about":
        return await about(update, ctx)

    if q.data == "img_menu":
        return await q.edit_message_text("🖼 *Image Converter*\nPick target format:",
                                         parse_mode="Markdown", reply_markup=image_menu())

    if q.data == "doc_menu":
        return await q.edit_message_text("📄 *Document Converter*\nPick target format:",
                                         parse_mode="Markdown", reply_markup=doc_menu())

    if q.data == "case_menu":
        return await q.edit_message_text("🔤 *Text Case Converter*\nPick a mode:",
                                         parse_mode="Markdown", reply_markup=case_menu())

    if q.data.startswith("img:"):
        fmt = q.data.split(":", 1)[1]
        USER_STATE[uid] = {"mode": "image", "target": fmt}
        return await q.edit_message_text(
            f"🖼 Send me an *image* and I'll convert it to *{fmt}*.",
            parse_mode="Markdown",
        )

    if q.data.startswith("doc:"):
        fmt = q.data.split(":", 1)[1]
        USER_STATE[uid] = {"mode": "doc", "target": fmt}
        return await q.edit_message_text(
            f"📄 Send me a *file* (TXT, MD, CSV, HTML, DOCX) or *text* and I'll convert it to *{fmt}*.",
            parse_mode="Markdown",
        )

    if q.data.startswith("case:"):
        mode = q.data.split(":", 1)[1]
        USER_STATE[uid] = {"mode": "case", "target": mode}
        return await q.edit_message_text(
            f"🔤 Send me *text* and I'll apply *{mode}* case.",
            parse_mode="Markdown",
        )

# ---------- Image upload ----------
async def handle_photo(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    state = USER_STATE.get(uid)

    if not state or state.get("mode") != "image":
        return await update.message.reply_text(
            "⚠️ First pick *Image Converter* and a target format.",
            parse_mode="Markdown",
            reply_markup=main_menu(),
        )

    target = state["target"]
    await update.message.reply_text(f"⏳ Converting to {target}...")

    photo = update.message.photo[-1] if update.message.photo else None
    document = update.message.document

    try:
        if photo:
            file = await photo.get_file()
            data = await file.download_as_bytearray()
        elif document and document.mime_type and document.mime_type.startswith("image/"):
            file = await document.get_file()
            data = await file.download_as_bytearray()
        else:
            return await update.message.reply_text("⚠️ Please send an image.")

        buf = convert_image(bytes(data), target)
        await update.message.reply_document(
            document=buf,
            filename=buf.name,
            caption=f"✅ Converted to *{target}*",
            parse_mode="Markdown",
        )
        USER_STATE.pop(uid, None)

    except Exception as e:
        logger.exception("Image conversion failed")
        await update.message.reply_text(f"❌ Conversion failed: {e}")

# ---------- Document upload ----------
async def handle_document(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    state = USER_STATE.get(uid)
    doc = update.message.document

    # If they sent an image file while in image mode
    if state and state.get("mode") == "image" and doc and doc.mime_type and doc.mime_type.startswith("image/"):
        return await handle_photo(update, ctx)

    if not state or state.get("mode") != "doc":
        return await update.message.reply_text(
            "⚠️ First pick *Document Converter* and a target format.",
            parse_mode="Markdown",
            reply_markup=main_menu(),
        )

    target = state["target"]
    await update.message.reply_text(f"⏳ Converting to {target}...")

    try:
        file = await doc.get_file()
        data = await file.download_as_bytearray()
        text = extract_text_from_upload(bytes(data), doc.file_name or "")
        buf = convert_document(text, target)
        await update.message.reply_document(
            document=buf,
            filename=buf.name,
            caption=f"✅ Converted to *{target}*",
            parse_mode="Markdown",
        )
        USER_STATE.pop(uid, None)
    except Exception as e:
        logger.exception("Document conversion failed")
        await update.message.reply_text(f"❌ Conversion failed: {e}")

# ---------- Text message ----------
async def handle_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    text = update.message.text
    if not text:
        return
    state = USER_STATE.get(uid)

    if not state:
        return await update.message.reply_text(
            "💡 Pick a tool from the menu, or send /menu.",
            reply_markup=main_menu(),
        )

    mode = state.get("mode")
    target = state.get("target")

    try:
        if mode == "case":
            result = apply_case(text, target)
            await update.message.reply_text(
                f"🔤 *Result:*\n`{result}`",
                parse_mode="Markdown",
            )
            USER_STATE.pop(uid, None)

        elif mode == "doc":
            buf = convert_document(text, target)
            await update.message.reply_document(
                document=buf,
                filename=buf.name,
                caption=f"✅ Converted to *{target}*",
                parse_mode="Markdown",
            )
            USER_STATE.pop(uid, None)

        else:
            await update.message.reply_text(
                "⚠️ Please send an image for image conversion.",
                reply_markup=main_menu(),
            )
            USER_STATE.pop(uid, None)

    except Exception as e:
        logger.exception("Text conversion failed")
        await update.message.reply_text(f"❌ Conversion failed: {e}")

async def error_handler(update: object, ctx: ContextTypes.DEFAULT_TYPE):
    logger.error("Exception:", exc_info=ctx.error)

# ---------- Main ----------
def main():
    if not TOKEN:
        logger.error("❌ BOT_TOKEN missing — set it in Railway Variables")
        time.sleep(30)
        raise SystemExit(1)

    logger.info(f"🚀 Starting SB24GZ Official Bot (token len={len(TOKEN)})")

    app = Application.builder().token(TOKEN).build()

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(CommandHandler("menu", menu_cmd))
    app.add_handler(CommandHandler("cancel", cancel_cmd))
    app.add_handler(CommandHandler("about", about))
    app.add_handler(CallbackQueryHandler(button_handler))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_error_handler(error_handler)

    logger.info("✅ Bot polling...")
    app.run_polling(drop_pending_updates=True, allowed_updates=Update.ALL_TYPES)

if __name__ == "__main__":
    main()
